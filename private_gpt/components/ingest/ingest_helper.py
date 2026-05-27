import logging
from pathlib import Path

from llama_index.core.readers import StringIterableReader
from llama_index.core.readers.base import BaseReader
from llama_index.core.readers.json import JSONReader
from llama_index.core.schema import Document

logger = logging.getLogger(__name__)


class SmartImageReader(BaseReader):
    """Custom image reader that extracts text using EasyOCR and gets visual descriptions."""

    def load_data(
        self, file_path: Path, extra_info: dict | None = None
    ) -> list[Document]:
        import easyocr
        from PIL import Image

        logger.info(f"Extracting text and details from image: {file_path}")
        
        # 1. OCR / Handwriting recognition using EasyOCR
        ocr_text = ""
        try:
            reader = easyocr.Reader(['en'], gpu=True) # It will auto-detect GPU/CUDA
            results = reader.readtext(str(file_path))
            ocr_text = "\n".join([res[1] for res in results])
        except Exception as e:
            logger.warning(f"EasyOCR failed: {e}")

        # 2. Image understanding / Visual description via OpenAI (if configured)
        description = ""
        import os
        openai_key = os.environ.get("OPENAI_API_KEY")
        if openai_key:
            try:
                import base64
                from openai import OpenAI

                # Encode image to base64
                with open(file_path, "rb") as image_file:
                    base64_image = base64.b64encode(image_file.read()).decode('utf-8')

                client = OpenAI(api_key=openai_key)
                response = client.chat.completions.create(
                    model="gpt-4o-mini",
                    messages=[
                        {
                            "role": "user",
                            "content": [
                                {"type": "text", "text": "Describe the contents of this image in detail, including any charts, diagrams, drawings, objects, or handwritten text. Provide a comprehensive summary that can be used for context in RAG search."},
                                {
                                    "type": "image_url",
                                    "image_url": {
                                        "url": f"data:image/jpeg;base64,{base64_image}"
                                    }
                                }
                            ]
                        }
                    ],
                    max_tokens=1000,
                )
                description = response.choices[0].message.content
            except Exception as e:
                logger.warning(f"OpenAI multimodal analysis failed: {e}")

        # Combine OCR and visual description
        full_text_parts = []
        if ocr_text:
            full_text_parts.append(f"--- OCR Extracted Text ---\n{ocr_text}")
        if description:
            full_text_parts.append(f"--- Visual Content Description ---\n{description}")

        full_text = "\n\n".join(full_text_parts) if full_text_parts else "Empty or unreadable image."
        
        return [Document(text=full_text, metadata=extra_info or {})]


class SmartPptxReader(BaseReader):
    """PowerPoint reader that extracts slide text efficiently without external model downloads."""

    def load_data(
        self, file_path: Path, extra_info: dict | None = None
    ) -> list[Document]:
        from pptx import Presentation

        logger.info(f"Extracting text from PowerPoint: {file_path}")
        presentation = Presentation(file_path)
        result = []
        for i, slide in enumerate(presentation.slides):
            slide_text_parts = []
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    slide_text_parts.append(shape.text.strip())
            if slide_text_parts:
                slide_text = "\n".join(slide_text_parts)
                result.append(f"Slide #{i + 1}:\n{slide_text}")
        
        full_text = "\n\n".join(result) if result else "Empty PowerPoint presentation."
        return [Document(text=full_text, metadata=extra_info or {})]


# Inspired by the `llama_index.core.readers.file.base` module
def _try_loading_included_file_formats() -> dict[str, type[BaseReader]]:
    try:
        from llama_index.readers.file.docs import (  # type: ignore
            DocxReader,
            HWPReader,
            PDFReader,
        )
        from llama_index.readers.file.epub import EpubReader  # type: ignore
        from llama_index.readers.file.ipynb import IPYNBReader  # type: ignore
        from llama_index.readers.file.markdown import MarkdownReader  # type: ignore
        from llama_index.readers.file.mbox import MboxReader  # type: ignore
        from llama_index.readers.file.tabular import (  # type: ignore
            PandasCSVReader,
            PandasExcelReader,
        )
        from llama_index.readers.file.video_audio import (  # type: ignore
            VideoAudioReader,
        )
    except ImportError as e:
        raise ImportError("`llama-index-readers-file` package not found") from e

    default_file_reader_cls: dict[str, type[BaseReader]] = {
        ".hwp": HWPReader,
        ".pdf": PDFReader,
        ".docx": DocxReader,
        ".pptx": SmartPptxReader,
        ".ppt": SmartPptxReader,
        ".pptm": SmartPptxReader,
        ".jpg": SmartImageReader,
        ".png": SmartImageReader,
        ".jpeg": SmartImageReader,
        ".mp3": VideoAudioReader,
        ".mp4": VideoAudioReader,
        ".csv": PandasCSVReader,
        ".xlsx": PandasExcelReader,
        ".xls": PandasExcelReader,
        ".epub": EpubReader,
        ".md": MarkdownReader,
        ".mbox": MboxReader,
        ".ipynb": IPYNBReader,
    }
    return default_file_reader_cls


# Patching the default file reader to support other file types
FILE_READER_CLS = _try_loading_included_file_formats()
FILE_READER_CLS.update(
    {
        ".json": JSONReader,
    }
)


class IngestionHelper:
    """Helper class to transform a file into a list of documents.

    This class should be used to transform a file into a list of documents.
    These methods are thread-safe (and multiprocessing-safe).
    """

    @staticmethod
    def transform_file_into_documents(
        file_name: str, file_data: Path
    ) -> list[Document]:
        documents = IngestionHelper._load_file_to_documents(file_name, file_data)
        for document in documents:
            document.metadata["file_name"] = file_name
        IngestionHelper._exclude_metadata(documents)
        return documents

    @staticmethod
    def _load_file_to_documents(file_name: str, file_data: Path) -> list[Document]:
        logger.debug("Transforming file_name=%s into documents", file_name)
        extension = Path(file_name).suffix
        reader_cls = FILE_READER_CLS.get(extension)
        if reader_cls is None:
            logger.debug(
                "No reader found for extension=%s, using default string reader",
                extension,
            )
            # Read as a plain text
            string_reader = StringIterableReader()
            return string_reader.load_data([file_data.read_text()])

        logger.debug("Specific reader found for extension=%s", extension)
        documents = reader_cls().load_data(file_data)

        # Sanitize NUL bytes in text which can't be stored in Postgres
        for i in range(len(documents)):
            documents[i].text = documents[i].text.replace("\u0000", "")

        return documents

    @staticmethod
    def _exclude_metadata(documents: list[Document]) -> None:
        logger.debug("Excluding metadata from count=%s documents", len(documents))
        for document in documents:
            document.metadata["doc_id"] = document.doc_id
            # We don't want the Embeddings search to receive this metadata
            document.excluded_embed_metadata_keys = ["doc_id"]
            # We don't want the LLM to receive these metadata in the context
            document.excluded_llm_metadata_keys = ["file_name", "doc_id", "page_label"]
