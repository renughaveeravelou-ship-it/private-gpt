import os
from pathlib import Path

# Create directory
test_dir = Path("test_docs")
test_dir.mkdir(exist_ok=True)

# 1. Generate Excel (.xlsx) using openpyxl
import openpyxl
wb = openpyxl.Workbook()
ws = wb.active
ws.title = "Sales Data"
ws.append(["Quarter", "Revenue", "Expenses"])
ws.append(["Q1", 10000, 6000])
ws.append(["Q2", 12000, 7000])
ws.append(["Q3", 15000, 8000])
ws.append(["Q4", 22000, 11000])
wb.save(test_dir / "test_sheet.xlsx")
print("Generated Excel file")

# 2. Generate PPTX (.pptx) using python-pptx
from pptx import Presentation
prs = Presentation()
slide = prs.slides.add_slide(prs.slide_layouts[1])
title = slide.shapes.title
subtitle = slide.placeholders[1]
title.text = "Introduction to Deep Learning"
subtitle.text = "Deep learning uses neural networks to learn representations from data. Convolutional Neural Networks (CNNs) are ideal for images."
prs.save(test_dir / "test_presentation.pptx")
print("Generated PPTX file")

# 3. Generate Image (.png) using PIL with text
from PIL import Image, ImageDraw, ImageFont
# Create a white image
img = Image.new("RGB", (600, 200), color="white")
d = ImageDraw.Draw(img)
# Draw text (using default font as we might not have external ttf easily located)
# Drawing some clear text
d.text((10, 10), "OCR Verification Note:", fill="black")
d.text((10, 50), "Important: The server database password is 'PGPT_SECURE_2026'.", fill="black")
d.text((10, 90), "Handwritten Style: Task A must be completed by June 5th.", fill="black")
img.save(test_dir / "test_image.png")
print("Generated Image file")
